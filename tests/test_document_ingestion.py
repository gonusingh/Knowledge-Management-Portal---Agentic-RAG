import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from qdrant_client import QdrantClient, models

import app.main as main_module
from app.agents.nodes import responder
from app.ingestion.chunking import MAX_CHUNK_CHARS, chunk_document
from app.ingestion.loaders import load_document
from app.security.access_control import DEMO_USERS
import app.services.retrieval.vector_search as vector_search


def make_text_pdf(text: str) -> bytes:
    """Build a minimal text PDF fixture without adding a test dependency."""
    escaped_text = (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = (
        f"BT /F1 10 Tf 36 750 Td ({escaped_text}) Tj ET"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        ),
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_number} 0 obj\n".encode("ascii"))
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(pdf)


class DocumentIngestionTests(unittest.TestCase):
    def test_pdf_resume_text_is_extracted_and_long_lines_are_chunked(self) -> None:
        resume_text = (
            "Aarush Sharma is a software engineer with Python, FastAPI, "
            "and retrieval experience. "
        ) * 40

        with TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "aarush_resume.pdf"
            pdf_path.write_bytes(make_text_pdf(resume_text))

            document = load_document(str(pdf_path))
            chunks = chunk_document(document)

        self.assertIn("Aarush Sharma", document["text"])
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk["text"]) <= MAX_CHUNK_CHARS for chunk in chunks))
        self.assertIn("retrieval experience", "\n".join(chunk["text"] for chunk in chunks))
        self.assertTrue(
            all(chunk["document_id"] == document["document_id"] for chunk in chunks)
        )

    def test_unreadable_pdf_reports_a_parsing_error(self) -> None:
        with TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "broken.pdf"
            pdf_path.write_bytes(b"not a PDF")

            with self.assertRaisesRegex(ValueError, "could not be parsed"):
                load_document(str(pdf_path))

    def test_textless_pdf_upload_returns_a_clear_client_error(self) -> None:
        original_overrides = dict(main_module.app.dependency_overrides)
        main_module.app.dependency_overrides[main_module.resolve_demo_user] = (
            lambda: DEMO_USERS["carol"]
        )

        try:
            with patch("app.main.get_document_metadata", return_value=None), patch(
                "app.main.upsert_chunks"
            ) as upsert:
                with TestClient(main_module.app) as client:
                    response = client.post(
                        "/upload",
                        files={
                            "file": (
                                "image_only.pdf",
                                make_text_pdf(""),
                                "application/pdf",
                            )
                        },
                        headers={"X-Demo-User": "carol"},
                    )

            self.assertEqual(response.status_code, 400)
            self.assertIn("OCR is not supported", response.json()["detail"])
            upsert.assert_not_called()
        finally:
            main_module.app.dependency_overrides = original_overrides

    def test_docx_table_text_is_available_to_retrieval(self) -> None:
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "aarush_resume.docx"
            document_file = Document()
            document_file.add_heading("Professional experience", level=1)
            table = document_file.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Candidate"
            table.cell(0, 1).text = "Aarush Sharma"
            document_file.save(docx_path)

            document = load_document(str(docx_path))
            chunks = chunk_document(document)

        self.assertIn("# Professional experience", document["text"])
        self.assertIn("Candidate | Aarush Sharma", document["text"])
        self.assertTrue(
            any("Aarush Sharma" in chunk["text"] for chunk in chunks)
        )

    def test_configured_chunk_overlap_does_not_cross_acl_boundary(self) -> None:
        from app.ingestion import chunking

        document = {
            "text": (
                "Public guidance " * 24
                + "\n\n"
                + "Restricted security procedure " * 24
            ),
            "source_file": "policy.txt",
            "type": "true_data",
            "document_id": "d" * 64,
            "classification": "RESTRICTED",
            "allowed_roles": ["administrator"],
            "tenant": "nimbuspay",
            "service": None,
            "effective_date": None,
            "version": "v1",
            "trusted_sections": [
                {
                    "start_paragraph": 0,
                    "end_paragraph": 1,
                    "classification": "PUBLIC",
                    "allowed_roles": ["viewer", "operator", "administrator"],
                },
                {
                    "start_paragraph": 1,
                    "end_paragraph": 2,
                    "classification": "RESTRICTED",
                    "allowed_roles": ["administrator"],
                },
            ],
        }
        with patch.object(chunking, "MAX_CHUNK_CHARS", 256), patch.object(
            chunking, "CHUNK_OVERLAP_CHARS", 32
        ):
            chunks = chunking.chunk_document(document)

        self.assertGreater(len(chunks), 2)
        self.assertTrue(
            all(
                chunk["classification"] == "PUBLIC"
                for chunk in chunks
                if "Public guidance" in chunk["text"]
                and "Restricted security procedure" not in chunk["text"]
            )
        )
        self.assertTrue(
            all(
                chunk["classification"] == "RESTRICTED"
                for chunk in chunks
                if "Restricted security procedure" in chunk["text"]
            )
        )

    def test_docx_tables_do_not_shift_trusted_section_paragraph_ranges(self) -> None:
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "sectioned.docx"
            document_file = Document()
            document_file.add_paragraph("Public section")
            table = document_file.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Public table detail"
            document_file.add_paragraph("Confidential section")
            document_file.save(docx_path)

            document = load_document(str(docx_path))
            document["trusted_sections"] = [
                {
                    "start_paragraph": 0,
                    "end_paragraph": 1,
                    "classification": "PUBLIC",
                    "allowed_roles": ["viewer"],
                },
                {
                    "start_paragraph": 1,
                    "end_paragraph": 2,
                    "classification": "CONFIDENTIAL",
                    "allowed_roles": ["administrator"],
                },
            ]
            chunks = chunk_document(document)

        self.assertEqual(len(chunks), 2)
        self.assertIn("Public table detail", chunks[0]["text"])
        self.assertEqual(chunks[0]["classification"], "PUBLIC")
        self.assertIn("Confidential section", chunks[1]["text"])
        self.assertEqual(chunks[1]["classification"], "CONFIDENTIAL")

    def test_pdf_upload_retrieval_and_answer_context_share_cv_content(self) -> None:
        resume_text = (
            "Aarush Sharma is a software engineer with Python and FastAPI "
            "experience. "
        ) * 30
        pdf_bytes = make_text_pdf(resume_text)
        original_overrides = dict(main_module.app.dependency_overrides)
        original_client = vector_search.client
        original_collection = vector_search.COLLECTION_NAME
        original_embedding = vector_search._get_embedding
        original_indexes = vector_search._payload_indexes_ready_for
        local_client = QdrantClient(":memory:")
        vector_search.client = local_client
        vector_search.COLLECTION_NAME = "resume_upload_test"
        vector_search._get_embedding = lambda _text: [
            1.0
        ] * vector_search.EMBEDDING_DIMENSIONS
        vector_search._payload_indexes_ready_for = None
        main_module.app.dependency_overrides[main_module.resolve_demo_user] = (
            lambda: DEMO_USERS["carol"]
        )

        try:
            with TestClient(main_module.app) as client:
                response = client.post(
                    "/upload",
                    files={"file": ("aarush_resume.pdf", pdf_bytes, "application/pdf")},
                    headers={"X-Demo-User": "carol"},
                )

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertGreater(payload["chunks_added"], 1)

            stored_points, _ = local_client.scroll(
                collection_name=vector_search.COLLECTION_NAME,
                scroll_filter=models.Filter(
                    must=[
                        models.FieldCondition(
                            key="document_id",
                            match=models.MatchValue(value=payload["document_id"]),
                        ),
                        models.FieldCondition(
                            key="tenant",
                            match=models.MatchValue(value=DEMO_USERS["carol"].tenant),
                        ),
                    ]
                ),
                limit=100,
                with_payload=True,
                with_vectors=True,
            )
            stale_payload = dict(stored_points[0].payload)
            stale_payload["chunk_id"] = f'{payload["document_id"]}_stale'
            stale_payload["chunk_index"] = payload["chunks_added"] + 3
            local_client.upsert(
                collection_name=vector_search.COLLECTION_NAME,
                points=[
                    models.PointStruct(
                        id=10000,
                        vector=[1.0] * vector_search.EMBEDDING_DIMENSIONS,
                        payload=stale_payload,
                    )
                ],
            )

            with TestClient(main_module.app) as client:
                refreshed_response = client.post(
                    "/upload",
                    files={"file": ("aarush_resume.pdf", pdf_bytes, "application/pdf")},
                    headers={"X-Demo-User": "carol"},
                )
            self.assertEqual(refreshed_response.status_code, 200, refreshed_response.text)
            self.assertTrue(refreshed_response.json()["reindexed"])

            refreshed_points, _ = local_client.scroll(
                collection_name=vector_search.COLLECTION_NAME,
                scroll_filter=models.Filter(
                    must=[
                        models.FieldCondition(
                            key="document_id",
                            match=models.MatchValue(value=payload["document_id"]),
                        ),
                        models.FieldCondition(
                            key="tenant",
                            match=models.MatchValue(value=DEMO_USERS["carol"].tenant),
                        ),
                    ]
                ),
                limit=100,
                with_payload=True,
                with_vectors=False,
            )
            self.assertEqual(len(refreshed_points), payload["chunks_added"])
            self.assertFalse(
                any(point.payload["chunk_id"].endswith("_stale") for point in refreshed_points)
            )

            retrieved = vector_search.search(
                query="Which candidate has FastAPI experience?",
                user=DEMO_USERS["carol"],
                top_k=5,
                document_id=payload["document_id"],
            )
            self.assertTrue(retrieved)
            self.assertTrue(
                any(
                    "Aarush Sharma" in chunk["text"]
                    and "FastAPI" in chunk["text"]
                    for chunk in retrieved
                )
            )

            state = {
                "user_query": "Which candidate has FastAPI experience?",
                "intent": "technical",
                "reranked_chunks": retrieved,
                "trace": [],
            }
            with patch.object(
                responder,
                "call_llm",
                return_value="Aarush Sharma has FastAPI experience.",
            ) as call_llm:
                answer = responder.responder_node(state)

            self.assertIn("Aarush Sharma", answer["final_answer"])
            context_chunks = call_llm.call_args.kwargs["context_chunks"]
            self.assertTrue(
                any("Aarush Sharma" in chunk["text"] for chunk in context_chunks)
            )
        finally:
            main_module.app.dependency_overrides = original_overrides
            vector_search.client = original_client
            vector_search.COLLECTION_NAME = original_collection
            vector_search._get_embedding = original_embedding
            vector_search._payload_indexes_ready_for = original_indexes
            local_client.close()


if __name__ == "__main__":
    unittest.main()
